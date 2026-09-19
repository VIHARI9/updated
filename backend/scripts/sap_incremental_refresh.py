from __future__ import annotations

import argparse
import json
import logging
import shutil
import sys
import time
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Callable

import pythoncom
import win32com.client
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


DAILY_DATE_HEADERS = {"date", "day", "productiondate", "postingdate", "reportdate"}
MONTH_HEADERS = {"month", "monthyear", "period", "date"}


@dataclass(frozen=True)
class Report:
    original: str
    temporary: str
    monthly: bool
    report_kind: str
    grid_column: str
    save_window: int


REPORTS = (
    Report("Daywise Data.xlsx", "Daywise Data_TEMP.xlsx", False, "daywise_data", "SALE_A_GRADE", 3),
    Report("Daywise MW Report.xlsx", "Daywise MW Report_TEMP.xlsx", False, "daywise_mw", "BEL_GRADE", 2),
    Report("Monthwise MW Report.xlsx", "Monthwise MW Report_TEMP.xlsx", True, "monthwise_mw", "A_GRADE", 1),
    Report("Monthwise Report.xlsx", "Monthwise Report_TEMP.xlsx", True, "monthwise_data", "SALE_A_GRADE", 1),
)


def normalize(value: object) -> str:
    return "".join(ch for ch in str(value or "").lower() if ch.isalnum())


def parse_period(value: object, monthly: bool, epoch) -> date | None:
    if value is None or value == "":
        return None
    parsed: date | None = None
    if isinstance(value, datetime):
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    elif isinstance(value, (int, float)):
        try:
            converted = from_excel(value, epoch)
            parsed = converted.date() if isinstance(converted, datetime) else converted
        except Exception:
            parsed = None
    else:
        text = str(value).strip()
        for fmt in (
            "%d-%b-%Y", "%d-%b-%y", "%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d",
            "%b-%Y", "%b-%y", "%B-%Y", "%B-%y", "%m-%Y", "%m/%Y",
        ):
            try:
                parsed = datetime.strptime(text, fmt).date()
                break
            except ValueError:
                pass
    if parsed and monthly:
        parsed = parsed.replace(day=1)
    return parsed


def find_data_sheet_and_date_column(workbook_path: Path, monthly: bool):
    wb = load_workbook(workbook_path, read_only=True, data_only=False)
    wanted = MONTH_HEADERS if monthly else DAILY_DATE_HEADERS
    best = None
    for ws in wb.worksheets:
        max_scan_row = min(ws.max_row, 25)
        max_scan_col = min(ws.max_column, 100)
        for row in range(1, max_scan_row + 1):
            for col in range(1, max_scan_col + 1):
                header = normalize(ws.cell(row, col).value)
                if header in wanted or (monthly and "month" in header) or (not monthly and "date" in header):
                    score = 100 if header in wanted else 50
                    candidate = (score, ws.title, row, col, wb.epoch)
                    if best is None or candidate[0] > best[0]:
                        best = candidate
    wb.close()
    if best is None:
        raise RuntimeError(f"No {'Month' if monthly else 'Date'} column found in {workbook_path.name}")
    return best[1], best[2], best[3], best[4]


def latest_period(workbook_path: Path, monthly: bool) -> date:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Missing workbook: {workbook_path}")
    sheet_name, header_row, date_column, epoch = find_data_sheet_and_date_column(workbook_path, monthly)
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    values = []
    for row in range(header_row + 1, ws.max_row + 1):
        parsed = parse_period(ws.cell(row, date_column).value, monthly, epoch)
        if parsed:
            values.append(parsed)
    wb.close()
    if not values:
        raise RuntimeError(f"No valid {'month' if monthly else 'date'} values in {workbook_path.name}")
    return max(values)


def connect_sap():
    try:
        sap_gui = win32com.client.GetObject("SAPGUI")
        application = sap_gui.GetScriptingEngine
        if application.Children.Count == 0:
            raise RuntimeError("No active SAP connection was found")
        connection = application.Children(0)
        if connection.Children.Count == 0:
            raise RuntimeError("No active SAP session was found")
        return application, connection, connection.Children(0)
    except Exception as exc:
        raise RuntimeError(f"SAP GUI connection failed: {exc}") from exc


def wait_for_object(session, object_id: str, timeout: int = 60):
    deadline = time.monotonic() + timeout
    last_error = None
    while time.monotonic() < deadline:
        try:
            return session.findById(object_id)
        except Exception as exc:
            last_error = exc
            time.sleep(0.25)
    raise TimeoutError(f"Timed out waiting for SAP object {object_id}: {last_error}")


def wait_for_file(path: Path, timeout: int = 120):
    deadline = time.monotonic() + timeout
    previous = -1
    stable = 0
    while time.monotonic() < deadline:
        if path.exists():
            size = path.stat().st_size
            if size > 0 and size == previous:
                stable += 1
                if stable >= 3:
                    try:
                        wb = load_workbook(path, read_only=True, data_only=True)
                        wb.close()
                        return
                    except Exception:
                        pass
            else:
                stable = 0
            previous = size
        time.sleep(0.5)
    raise TimeoutError(f"Export was not created or remained locked: {path}")


def set_period(session, start: date, end: date):
    low = wait_for_object(session, "wnd[0]/usr/ctxtS_DATE-LOW")
    high = wait_for_object(session, "wnd[0]/usr/ctxtS_DATE-HIGH")
    low.text = start.strftime("%d.%m.%Y")
    high.text = end.strftime("%d.%m.%Y")
    high.setFocus()
    session.findById("wnd[0]").sendVKey(0)
    time.sleep(0.5)


def wait_until_sap_ready(session, timeout: int = 120):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        try:
            if not session.Busy:
                return
        except Exception:
            pass
        time.sleep(0.25)
    raise TimeoutError("SAP remained busy for more than 120 seconds")


def sap_status_text(session) -> str:
    try:
        return str(session.findById("wnd[0]/sbar").text or "").strip()
    except Exception:
        return ""


def current_transaction(session) -> str:
    try:
        return str(session.Info.Transaction or "").strip()
    except Exception:
        return ""


def open_transaction(session, command: str, start: date, end: date):
    # /n guarantees that SAP leaves the current transaction and opens a clean
    # ZCELLREPORT selection screen. A bare transaction code can be ignored when
    # the current screen has pending state or an active report result.
    transaction = "/nzcellreport"
    command_field = wait_for_object(session, "wnd[0]/tbar[0]/okcd", 30)
    command_field.text = transaction
    session.findById("wnd[0]").sendVKey(0)
    wait_until_sap_ready(session)

    try:
        wait_for_object(session, "wnd[0]/usr/ctxtS_DATE-LOW", 30)
        wait_for_object(session, "wnd[0]/usr/ctxtS_DATE-HIGH", 30)
    except TimeoutError as exc:
        status = sap_status_text(session)
        active_transaction = current_transaction(session)
        details = (
            f"ZCELLREPORT selection screen did not open. "
            f"Current transaction: {active_transaction or 'unknown'}. "
            f"SAP status: {status or 'no status message'}."
        )
        raise RuntimeError(details) from exc

    set_period(session, start, end)


def press(session, object_id: str, timeout: int = 60):
    wait_for_object(session, object_id, timeout).press()
    time.sleep(0.5)


def export_grid(session, data_dir: Path, report: Report):
    grid = wait_for_object(session, "wnd[0]/usr/cntlGRID1/shellcont/shell", 120)
    grid.currentCellColumn = report.grid_column
    grid.selectedRows = "0"
    grid.contextMenu()
    grid.selectContextMenuItem("&XXL")
    press(session, "wnd[1]/tbar[0]/btn[0]", 60)

    if report.save_window == 3:
        session.findById("wnd[1]").sendVKey(4)
        time.sleep(0.5)
        session.findById("wnd[2]").sendVKey(4)
    elif report.save_window == 2:
        session.findById("wnd[1]").sendVKey(4)

    save_window = f"wnd[{report.save_window}]"
    wait_for_object(session, save_window + "/usr/ctxtDY_PATH", 60).text = str(data_dir)
    session.findById(save_window + "/usr/ctxtDY_FILENAME").text = report.temporary
    session.findById(save_window + "/tbar[0]/btn[11]").press()
    time.sleep(1)

    for window_number in range(report.save_window - 1, 0, -1):
        try:
            session.findById(f"wnd[{window_number}]/tbar[0]/btn[11]").press()
            time.sleep(0.3)
        except Exception:
            pass
    wait_for_file(data_dir / report.temporary)


def export_report(session, data_dir: Path, report: Report, start: date, end: date):
    if report.report_kind == "daywise_data":
        open_transaction(session, "zcellreport", start, end)
        press(session, "wnd[0]/tbar[1]/btn[8]")
        radio = wait_for_object(session, "wnd[0]/usr/radR_BUT4")
        radio.Select(); radio.SetFocus()
        press(session, "wnd[0]/tbar[1]/btn[8]")
    elif report.report_kind == "daywise_mw":
        open_transaction(session, "/nzcellreport", start, end)
        press(session, "wnd[0]/tbar[1]/btn[8]")
        press(session, "wnd[0]/tbar[1]/btn[8]")
    elif report.report_kind == "monthwise_mw":
        open_transaction(session, "/nzcellreport", start, end)
        radio = wait_for_object(session, "wnd[0]/usr/radR_BUT2")
        radio.Select(); radio.SetFocus()
        press(session, "wnd[0]/tbar[1]/btn[8]")
        press(session, "wnd[0]/tbar[1]/btn[8]")
    else:
        open_transaction(session, "/nzcellreport", start, end)
        radio = wait_for_object(session, "wnd[0]/usr/radR_BUT2")
        radio.Select(); radio.SetFocus()
        press(session, "wnd[0]/tbar[1]/btn[8]")
        radio = wait_for_object(session, "wnd[0]/usr/radR_BUT4")
        radio.Select(); radio.SetFocus()
        press(session, "wnd[0]/tbar[1]/btn[8]")
    export_grid(session, data_dir, report)


def copy_cell(source, target):
    target.value = source.value
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    if source.font:
        target.font = copy(source.font)
    if source.fill:
        target.fill = copy(source.fill)
    if source.border:
        target.border = copy(source.border)
    if source.alignment:
        target.alignment = copy(source.alignment)
    if source.protection:
        target.protection = copy(source.protection)


def merge_report(original_path: Path, temp_path: Path, boundary: date, monthly: bool):
    original_wb = load_workbook(original_path)
    temp_wb = load_workbook(temp_path, data_only=False)
    original_sheet, original_header, original_col, _ = find_data_sheet_and_date_column(original_path, monthly)
    temp_sheet, temp_header, temp_col, _ = find_data_sheet_and_date_column(temp_path, monthly)
    original_ws = original_wb[original_sheet]
    temp_ws = temp_wb[temp_sheet]

    delete_rows = []
    for row in range(original_header + 1, original_ws.max_row + 1):
        parsed = parse_period(original_ws.cell(row, original_col).value, monthly, original_wb.epoch)
        if parsed and parsed >= boundary:
            delete_rows.append(row)
    for row in reversed(delete_rows):
        original_ws.delete_rows(row, 1)

    source_rows = []
    for row in range(temp_header + 1, temp_ws.max_row + 1):
        parsed = parse_period(temp_ws.cell(row, temp_col).value, monthly, temp_wb.epoch)
        if parsed and parsed >= boundary:
            source_rows.append(row)
    if not source_rows:
        original_wb.close(); temp_wb.close()
        raise RuntimeError(f"{temp_path.name} contains no rows on or after {boundary}")

    target_row = original_ws.max_row + 1
    for source_row in source_rows:
        for col in range(1, temp_ws.max_column + 1):
            copy_cell(temp_ws.cell(source_row, col), original_ws.cell(target_row, col))
        target_row += 1

    validation_path = original_path.with_name(original_path.stem + "_VALIDATED.xlsx")
    original_wb.save(validation_path)
    original_wb.close(); temp_wb.close()

    check = load_workbook(validation_path, read_only=True, data_only=True)
    check.close()
    return validation_path


def close_exported_excel_workbooks(data_dir: Path):
    try:
        excel = win32com.client.GetActiveObject("Excel.Application")
    except Exception:
        return
    expected = {(data_dir / report.temporary).resolve().as_posix().lower() for report in REPORTS}
    for index in range(excel.Workbooks.Count, 0, -1):
        workbook = excel.Workbooks(index)
        try:
            full_name = Path(workbook.FullName).resolve().as_posix().lower()
            if full_name in expected:
                workbook.Close(False)
        except Exception:
            pass


def run_refresh(data_dir: Path, progress: Callable[[int, str], None] | None = None):
    progress = progress or (lambda percent, message: None)
    data_dir = data_dir.resolve()
    data_dir.mkdir(parents=True, exist_ok=True)
    log_path = data_dir / "SAP Python Refresh Log.txt"
    logging.basicConfig(filename=log_path, level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s", force=True)
    started = time.monotonic()
    logging.info("Refresh started")

    boundaries = {}
    for index, report in enumerate(REPORTS):
        boundaries[report.original] = latest_period(data_dir / report.original, report.monthly)
        logging.info("%s last period: %s", report.original, boundaries[report.original])
        progress(5 + index * 3, f"Read {report.original}: {boundaries[report.original]}")
        temp = data_dir / report.temporary
        if temp.exists():
            temp.unlink()

    pythoncom.CoInitialize()
    try:
        application, connection, session = connect_sap()
        session.findById("wnd[0]").maximize()
        for index, report in enumerate(REPORTS):
            start = boundaries[report.original]
            progress(20 + index * 12, f"Exporting {report.original} from {start}")
            export_report(session, data_dir, report, start, date.today())
            logging.info("Exported %s from %s through %s", report.temporary, start, date.today())

        close_exported_excel_workbooks(data_dir)

        validated = []
        for index, report in enumerate(REPORTS):
            progress(70 + index * 5, f"Validating and merging {report.original}")
            validated_path = merge_report(
                data_dir / report.original,
                data_dir / report.temporary,
                boundaries[report.original],
                report.monthly,
            )
            validated.append((report, validated_path))

        backup_dir = data_dir / "refresh_backup"
        backup_dir.mkdir(exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backups = []
        try:
            for report, validated_path in validated:
                original = data_dir / report.original
                backup = backup_dir / f"{original.stem}_{timestamp}{original.suffix}"
                shutil.copy2(original, backup)
                backups.append((original, backup))
                validated_path.replace(original)
            for report in REPORTS:
                temp = data_dir / report.temporary
                if temp.exists():
                    temp.unlink()
        except Exception:
            for original, backup in backups:
                if backup.exists():
                    shutil.copy2(backup, original)
            raise

        try:
            from app.efficiency_service import build_efficiency_json
            build_efficiency_json(force=True)
            logging.info("Rebuilt sap_efficiency_daily.json")
        except Exception as exc:
            logging.warning("Report refresh succeeded, but efficiency JSON rebuild failed: %s", exc)

        duration = round(time.monotonic() - started, 1)
        message = f"SAP refresh completed in {duration} seconds"
        logging.info(message)
        progress(100, message)
        return {"status": "completed", "message": message, "log": str(log_path)}
    except Exception:
        logging.exception("Refresh failed")
        raise
    finally:
        pythoncom.CoUninitialize()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("data_dir", type=Path)
    args = parser.parse_args()
    result = run_refresh(args.data_dir, lambda percent, message: print(f"[{percent}%] {message}", flush=True))
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
